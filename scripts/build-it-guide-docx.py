#!/usr/bin/env python3
"""Builds the operational IT handover Word from the maintained Markdown guides."""
from __future__ import annotations

import re
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "GUIA_IMPLEMENTACAO_IT_START_CAMPUS.docx"
SOURCES = [
    ROOT / "docs" / "GUIA_IMPLEMENTACAO_IT_START_CAMPUS.md",
    ROOT / "docs" / "INTEGRACAO_BMS_FASE_2.md",
]

TEAL = "0A3638"
GREEN = "00A859"
INK = "18212F"
MUTED = "5E6C76"
PALE = "E9F5EF"
GRID = "D3DFDA"


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def border(cell, color: str = GRID) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), color)
        borders.append(node)
    tc_pr.append(borders)


def set_cell_text(cell, value: str, bold: bool = False, color: str = INK, size: int = 8) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(value)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border(cell)


def add_page_field(paragraph) -> None:
    paragraph.add_run("Página ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def set_default_fonts(document: Document) -> None:
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].font.color.rgb = RGBColor.from_string(INK)
    for level, size in (("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 10.5)):
        style = styles[level]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(TEAL)
        style.paragraph_format.space_before = Pt(16 if level == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(6)


def add_header_footer(document: Document) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.text = "STAND  /  START CAMPUS     GUIA DE IMPLEMENTAÇÃO IT"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(GREEN)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Uso interno · Informação operacional controlada · ").font.size = Pt(8)
    add_page_field(footer)
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)


def add_cover(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    for _ in range(3):
        document.add_paragraph()
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("START CAMPUS  ·  STAND")
    run.bold = True
    run.font.name = "Aptos"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(GREEN)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(20)
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("Guia de Implementação\npara a Equipa de IT")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(30)
    run.font.color.rgb = RGBColor.from_string(TEAL)
    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(28)
    run = sub.add_run("STAND — Plataforma de Governação Ambiental\nInstalação, segurança, integrações e Operação NEST / BMS")
    run.font.name = "Aptos"
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    band = document.add_table(rows=1, cols=3)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = [("VERSÃO", "Main / release validada"), ("AMBIENTE", "Staging → Produção"), ("SEGURANÇA", "Segredos fora do Git")]
    for index, (label, value) in enumerate(values):
        cell = band.cell(0, index)
        shade(cell, PALE)
        set_cell_text(cell, f"{label}\n{value}", bold=False, color=TEAL, size=9)
    document.add_paragraph()
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(22)
    run = note.add_run("Documento de passagem para IT. Não contém palavras-passe, chaves, tokens ou URLs privadas.")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    document.add_page_break()


def inline_text(value: str) -> str:
    value = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", value)
    value = value.replace("**", "").replace("`", "")
    return value.strip()


def add_table(document: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    headers = rows[0]
    body = rows[2:] if len(rows) > 2 and all(set(cell) <= {"-", ":"} for cell in rows[1]) else rows[1:]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, TEAL)
        set_cell_text(cell, inline_text(value), bold=True, color="FFFFFF", size=8)
    for row in body:
        cells = table.add_row().cells
        for index in range(len(headers)):
            text = inline_text(row[index]) if index < len(row) else ""
            set_cell_text(cells[index], text, size=8)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code(document: Document, code: list[str]) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, "102C2D")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("\n".join(code))
    run.font.name = "Consolas"
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor.from_string("E8FFF3")
    border(cell, "2E5554")


def add_markdown(document: Document, markdown: str, title_override: str | None = None) -> None:
    lines = markdown.splitlines()
    index = 0
    if title_override:
        document.add_heading(title_override, level=1)
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith("```"):
            code: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                code.append(lines[index])
                index += 1
            add_code(document, code)
            index += 1
            continue
        if line.startswith("|"):
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            add_table(document, rows)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            depth = min(len(heading.group(1)), 3)
            title = inline_text(heading.group(2))
            if depth == 1 and title_override:
                index += 1
                continue
            document.add_heading(title, level=depth)
            index += 1
            continue
        if line.startswith("> "):
            p = document.add_paragraph(style="Quote")
            p.add_run(inline_text(line[2:]))
            index += 1
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", line)
        numbered = re.match(r"^\d+[.)]\s+(.+)$", line)
        if bullet or numbered:
            style = "List Bullet" if bullet else "List Number"
            p = document.add_paragraph(style=style)
            p.add_run(inline_text((bullet or numbered).group(1)))
            index += 1
            continue
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.add_run(inline_text(line))
        index += 1


def main() -> None:
    document = Document()
    set_default_fonts(document)
    add_header_footer(document)
    add_cover(document)
    for source_index, source in enumerate(SOURCES):
        if source_index:
            document.add_page_break()
        markdown = source.read_text(encoding="utf-8")
        add_markdown(document, markdown, "Integração BMS — Fase 2 — SIN01 / NEST" if source.name.startswith("INTEGRACAO") else None)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
